# 천재이 중1 7과 업로드현황 xlsx (examples/YBM박8/make_xlsx.py 구조: 요약 · 파일 목록)
import os, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
OUT = sys.argv[1]; dst = sys.argv[2]
parts = [('1. 단어', '00. 단어'), ('2. 대화', '01. 대화'), ('3. 문법', '02. 문법'), ('4. 본문', '03. 본문')]
base = 'G:\\내 드라이브\\0. JLS수업\\3. 내신\\중1\\천재이\\00 공통자료\\'
def kind(f):
    if '암기지' in f: return '공용'
    return '교사용/답지' if ('답지' in f or '교사용' in f) else '학생용/문제'
wb = Workbook(); s = wb.active; s.title = '요약'
B = Font(bold=True); H = PatternFill('solid', fgColor='DDEBF7')
s.append(['천재이 7과 (중1 · 초당중) — 구글드라이브 업로드 현황']); s['A1'].font = Font(bold=True, size=14)
s.append(["드라이브 경로: 0. JLS수업 > 3. 내신 > 중1 > 천재이 > 00 공통자료 > … · 생성 2026-10-13 (클라우드 세션) · 개수는 '파일 목록' 시트에서 자동 계산"])
s.append([])
s.append(['파트', '드라이브 위치', '폴더 열기', '컴퓨터 위치', '개수'])
for c in s[4]: c.font = B; c.fill = H
for p, f in parts:
    r = s.max_row + 1
    s.append([p, f'00 공통자료 > {f} > 7과', '', base + f + '\\7과', f"=COUNTIF('파일 목록'!B:B,A{r})"])
s.append(['합계', '', '', '', '=SUM(E5:E8)']); s['A9'].font = B
s.append([])
s.append(['메모']); s['A11'].font = B
for m in ['드라이브 링크는 커넥터 있는 세션에서 채움 (폴더 열기·드라이브 링크 칸 비워 둠).',
          '중1 천재(이상기) Lesson 7 Tasty World (2022 개정) · 2026 2학기 기말고사 대비 · 초당중 1학년.',
          '문법 2개는 새 문법: 감각동사(_도구\\david\\build_G-감각동사_중1.js) · 비교급(build_G-비교급_중1.js).',
          '드라이브 중1 폴더 구조는 중3과 같은 방식(천재이 > 00 공통자료 > 파트 > 7과) — 원본 HWP도 같은 00 공통자료 폴더에 있음.']:
    s.append([m])
for col, w in zip('ABCDE', [12, 32, 12, 74, 8]): s.column_dimensions[col].width = w
t = wb.create_sheet('파일 목록'); t.append(['No', '파트', '파일명', '구분', '드라이브 링크'])
for c in t[1]: c.font = B; c.fill = H
n = 0
for p, f in parts:
    for fn in sorted(os.listdir(f'{OUT}/{f}/7과')):
        if fn.endswith('.docx'): n += 1; t.append([n, p, fn, kind(fn), ''])
for col, w in zip('ABCDE', [6, 10, 60, 14, 40]): t.column_dimensions[col].width = w
wb.save(dst); print(n)
