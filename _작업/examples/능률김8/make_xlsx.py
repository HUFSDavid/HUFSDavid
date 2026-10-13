# 능률김 8과 업로드현황 xlsx (examples/YBM박8/make_xlsx.py 기반 · 요약 · 파일 목록) — 목표 문법 없음 → 문법 파트 없이 3파트
import os, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
OUT=sys.argv[2] if len(sys.argv)>2 else '/tmp/david_w/out/능률김/8과'; dst=sys.argv[1]
parts=[('1. 단어','00. 단어'),('2. 대화','01. 대화'),('3. 본문','03. 본문')]
base='G:\\내 드라이브\\0. JLS수업\\★내신\\중3\\능률김\\00 공통자료\\'
def kind(f):
    if '암기지' in f: return '공용'
    return '교사용/답지' if ('답지' in f or '교사용' in f) else '학생용/문제'
wb=Workbook(); s=wb.active; s.title='요약'
B=Font(bold=True); H=PatternFill('solid',fgColor='DDEBF7')
s.append(['능률김 8과 (대지·어정중) — 구글드라이브 업로드 현황']); s['A1'].font=Font(bold=True,size=14)
s.append(["드라이브 경로: 0. JLS수업 > ★내신 > 중3 > 능률김 > 00 공통자료 > … · 재제작 2026-10-13 (클라우드 세션 — zip 전달) · 개수는 '파일 목록' 시트에서 자동 계산"])
s.append([])
s.append(['파트','드라이브 위치','폴더 열기','컴퓨터 위치','개수'])
for c in s[4]: c.font=B; c.fill=H
for p,f in parts:
    r=s.max_row+1
    s.append([p,f'00 공통자료 > {f} > 8과','',base+f+'\\8과',f'=COUNTIF(\'파일 목록\'!B:B,A{r})'])
s.append(['합계','','','','=SUM(E5:E7)']); s['A8'].font=B
s.append([])
s.append(['메모'])
s['A10'].font=B
for m in ['드라이브 링크는 커넥터 있는 세션에서 채움 (폴더 열기·드라이브 링크 칸 비워 둠).',
          '드라이브 8과 폴더 안의 기존 하위 폴더는 그대로 두고, 새 파일은 8과 폴더 바로 아래에 올림.',
          '8과는 Read On 단원(Listen & Talk 없음) → 대화 파트는 단원 안 대화 형식 지문 4개(Before You Read B · After You Read B · Project Work 1 Scene 2·5), 목표 문법 없음 → 문법 파트 없음(02. 문법 폴더 없음) · 총 27개.',
          '재제작(2026-10-13): 원래 결과 zip이 드라이브에 올라가지 않아 같은 방식으로 다시 만듦.']:
    s.append([m])
for col,w in zip('ABCDE',[12,32,12,70,8]): s.column_dimensions[col].width=w
t=wb.create_sheet('파일 목록'); t.append(['No','파트','파일명','구분','드라이브 링크'])
for c in t[1]: c.font=B; c.fill=H
n=0
for p,f in parts:
    for fn in sorted(os.listdir(f'{OUT}/{f}')):
        if fn.endswith('.docx'): n+=1; t.append([n,p,fn,kind(fn),''])
for col,w in zip('ABCDE',[6,10,60,14,40]): t.column_dimensions[col].width=w
wb.save(dst); print(n)
