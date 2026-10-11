import openpyxl,os,sys
from openpyxl.styles import Font,PatternFill,Alignment
OUT=sys.argv[1]; ROOT='/tmp/david_w/out/비상김/8과'
parts=[('1. 단어','00. 단어'),('2. 대화','01. 대화'),('3. 문법','02. 문법'),('4. 본문','03. 본문')]
order={'1. 단어':['0단계','1단계','2단계','3단계','4단계','5단계']}
wb=openpyxl.Workbook(); ws=wb.active; ws.title='요약'; fl=wb.create_sheet('파일 목록')
B=Font(bold=True); H=PatternFill('solid',fgColor='1F3864'); HF=Font(bold=True,color='FFFFFF')
ws['A1']='비상김 8과 (동백중) — 구글드라이브 업로드 현황'; ws['A1'].font=Font(bold=True,size=14)
ws['A2']="드라이브 경로: 0. JLS수업 > ★내신 > 중3 > 비상김 > 00 공통자료 > … · 생성 2026-10-11 (클라우드 세션 — zip 전달) · 개수는 '파일 목록' 시트에서 자동 계산"
ws.append([]); ws.append(['파트','드라이브 위치','폴더 열기','컴퓨터 위치','개수'])
for c in ws[4]: c.font=HF; c.fill=H
fl.append(['No','파트','파일명','구분','드라이브 링크'])
for c in fl[1]: c.font=HF; c.fill=H
n=0
def kind(f):
    if '학생용' in f or '_문제' in f: return '학생용/문제'
    if '교사용' in f or '_답지' in f: return '교사용/답지'
    return '공용'
for i,(pt,fd) in enumerate(parts):
    files=sorted(x for x in os.listdir(f'{ROOT}/{fd}') if x.endswith('.docx'))
    def key(f):
        k=0 if '문제' in f or '학생용' in f else 1
        return (f.replace('문제','').replace('답지','').replace('학생용','').replace('교사용',''),k)
    files.sort(key=key)
    for f in files:
        n+=1; fl.append([n,pt,f,kind(f),''])
    r=5+i
    ws.append([pt,f'00 공통자료 > {fd} > 8과','',f'G:\\내 드라이브\\0. JLS수업\\★내신\\중3\\비상김\\00 공통자료\\{fd}\\8과',f'=COUNTIF(\'파일 목록\'!B:B,A{r})'])
ws.append(['합계','','','','=SUM(E5:E8)']); ws['A9'].font=B
ws.append([]); ws.append(['메모'])
for m in ['드라이브 링크는 커넥터 있는 세션에서 채움 (폴더 열기·드라이브 링크 칸 비워 둠).',
          '드라이브 8과 폴더 안의 기존 하위 폴더는 그대로 두고, 새 파일은 8과 폴더 바로 아래에 올림.',
          '가정법 과거는 공통 문법(동아윤8) 스크립트를 COVER·PREFIX만 바꿔 재실행. with + 명사 + 분사는 새 공통 문법(_도구\\david\\build_G-with분사.js).',
          '다음 대기 세트: YBM박 7과 (죽전중 / 관계부사 how, the+비교급)']: ws.append([m])
for col,w in zip('ABCDE',[12,32,12,70,8]): ws.column_dimensions[col].width=w
for col,w in zip('ABCDE',[6,10,70,14,30]): fl.column_dimensions[col].width=w
wb.save(OUT); print('files',n)
